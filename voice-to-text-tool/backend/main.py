#!/usr/bin/env python3
"""
语音转文字工具 - FastAPI 后端
支持 ASR 转写、LLM 重点提取、多格式导出
"""

import os
import json
import asyncio
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, List
from contextlib import asynccontextmanager

from fastapi import FastAPI, File, UploadFile, Form, HTTPException, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
from dotenv import load_dotenv

# 加载环境变量
load_dotenv()

# 创建导出目录
EXPORT_DIR = Path("../exports")
EXPORT_DIR.mkdir(exist_ok=True)
UPLOAD_DIR = Path("../uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

# 全局 ASR 模型（延迟加载）
asr_model = None

# ============ Pydantic 模型 ============

class TranscriptionResult(BaseModel):
    task_id: str
    status: str  # pending, processing, completed, failed
    text: Optional[str] = None
    segments: Optional[List[dict]] = None
    error: Optional[str] = None
    created_at: str
    completed_at: Optional[str] = None

class SummaryRequest(BaseModel):
    task_id: str
    summary_type: str = "key_points"  # key_points, abstract, chapters

class SummaryResult(BaseModel):
    task_id: str
    summary: str
    summary_type: str

class ExportRequest(BaseModel):
    task_id: str
    format: str  # txt, srt, json, docx
    include_summary: bool = False

# ============ ASR 服务 ============

def get_asr_model():
    """懒加载 ASR 模型"""
    global asr_model
    if asr_model is None:
        print("🔄 正在加载 ASR 模型，首次启动可能需要几分钟...")
        try:
            from funasr import AutoModel
            # 使用 paraformer-zh，支持中文，带标点和VAD
            asr_model = AutoModel(
                model="paraformer-zh",
                model_revision="v2.0.4",
                vad_model="fsmn-vad",
                vad_model_revision="v2.0.4",
                punc_model="ct-punc-c",
                punc_model_revision="v2.0.4",
                device="cpu",  # 可根据环境改为 cuda
            )
            print("✅ ASR 模型加载完成")
        except Exception as e:
            print(f"❌ ASR 模型加载失败: {e}")
            raise
    return asr_model

async def process_audio_async(audio_path: str, task_id: str):
    """异步处理音频文件"""
    try:
        model = get_asr_model()
        
        # 更新任务状态
        update_task_status(task_id, "processing")
        
        # 执行 ASR（在线程池中运行，避免阻塞）
        loop = asyncio.get_event_loop()
        result = await loop.run_in_executor(
            None, 
            lambda: model.generate(input=audio_path, batch_size_s=300)
        )
        
        # 解析结果
        if result and len(result) > 0:
            text = result[0].get('text', '')
            
            # 保存结果
            task_result = {
                "task_id": task_id,
                "status": "completed",
                "text": text,
                "segments": result[0].get('sentence_info', []),
                "created_at": get_task(task_id).get("created_at"),
                "completed_at": datetime.now().isoformat()
            }
            save_task_result(task_id, task_result)
        else:
            raise Exception("ASR 返回空结果")
            
    except Exception as e:
        update_task_status(task_id, "failed", str(e))
        print(f"❌ 处理任务 {task_id} 失败: {e}")

# ============ LLM 服务 ============

async def generate_summary(text: str, summary_type: str = "key_points") -> str:
    """调用 LLM 生成摘要"""
    api_key = os.getenv("KIMI_API_KEY") or os.getenv("LLM_API_KEY")
    
    if not api_key:
        # 如果没有配置API Key，使用本地简单摘要
        return generate_local_summary(text, summary_type)
    
    try:
        import httpx
        
        prompts = {
            "key_points": "请从以下文本中提取5-8个关键要点，用 bullet points 形式列出：",
            "abstract": "请为以下文本生成一段200字左右的摘要：",
            "chapters": "请为以下内容生成章节大纲和时间戳："
        }
        
        prompt = prompts.get(summary_type, prompts["key_points"])
        
        async with httpx.AsyncClient(timeout=60.0) as client:
            response = await client.post(
                "https://api.moonshot.cn/v1/chat/completions",
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json"
                },
                json={
                    "model": os.getenv("LLM_MODEL", "kimi-k2.5"),
                    "messages": [
                        {"role": "system", "content": "你是一个专业的内容分析助手，擅长从语音转录文本中提取关键信息。"},
                        {"role": "user", "content": f"{prompt}\n\n{text[:8000]}"}  # 限制长度
                    ],
                    "temperature": 0.3
                }
            )
            
            if response.status_code == 200:
                data = response.json()
                return data["choices"][0]["message"]["content"]
            else:
                print(f"LLM API 错误: {response.status_code} - {response.text}")
                return generate_local_summary(text, summary_type)
                
    except Exception as e:
        print(f"LLM 调用失败: {e}")
        return generate_local_summary(text, summary_type)

def generate_local_summary(text: str, summary_type: str) -> str:
    """本地简单摘要（备用方案）"""
    sentences = text.split("。")
    
    if summary_type == "key_points":
        # 提取前8个句子作为要点
        key_points = [f"• {s.strip()}" for s in sentences[:8] if len(s.strip()) > 10]
        return "\n".join(key_points) if key_points else "• " + text[:500]
    
    elif summary_type == "abstract":
        # 取前3个句子作为摘要
        return "。".join(sentences[:3]) + "。" if len(sentences) > 1 else text[:500]
    
    else:
        return text[:1000]

# ============ 文件导出 ============

def export_to_txt(task_data: dict, include_summary: bool = False) -> str:
    """导出为 TXT 格式"""
    content = []
    content.append("=" * 50)
    content.append("语音转文字结果")
    content.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    content.append("=" * 50)
    content.append("")
    
    if include_summary and task_data.get("summary"):
        content.append("【内容摘要】")
        content.append(task_data["summary"])
        content.append("")
        content.append("=" * 50)
        content.append("")
    
    content.append("【转写内容】")
    content.append(task_data.get("text", ""))
    
    return "\n".join(content)

def export_to_srt(task_data: dict) -> str:
    """导出为 SRT 字幕格式"""
    segments = task_data.get("segments", [])
    if not segments:
        # 如果没有时间戳信息，生成单条字幕
        text = task_data.get("text", "")
        return f"1\n00:00:00,000 --> 00:00:30,000\n{text}\n"
    
    srt_lines = []
    for i, seg in enumerate(segments, 1):
        start_ms = seg.get("start", 0)
        end_ms = seg.get("end", 0)
        text = seg.get("text", "").strip()
        
        # 转换时间格式
        def ms_to_srt(ms):
            hours = ms // 3600000
            minutes = (ms % 3600000) // 60000
            seconds = (ms % 60000) // 1000
            millis = ms % 1000
            return f"{hours:02d}:{minutes:02d}:{seconds:02d},{millis:03d}"
        
        srt_lines.append(f"{i}")
        srt_lines.append(f"{ms_to_srt(start_ms)} --> {ms_to_srt(end_ms)}")
        srt_lines.append(text)
        srt_lines.append("")
    
    return "\n".join(srt_lines)

def export_to_docx(task_data: dict, include_summary: bool = False) -> str:
    """导出为 Word 文档"""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    
    # 标题
    title = doc.add_heading('语音转文字结果', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 元信息
    meta = doc.add_paragraph()
    meta.add_run(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n").italic = True
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # 空行
    
    # 摘要
    if include_summary and task_data.get("summary"):
        doc.add_heading('内容摘要', level=1)
        doc.add_paragraph(task_data["summary"])
        doc.add_page_break()
    
    # 转写内容
    doc.add_heading('转写内容', level=1)
    doc.add_paragraph(task_data.get("text", ""))
    
    # 保存
    output_path = EXPORT_DIR / f"{task_data['task_id']}.docx"
    doc.save(output_path)
    return str(output_path)

# ============ 任务管理 ============

tasks_db = {}

def create_task() -> str:
    """创建新任务"""
    task_id = str(uuid.uuid4())[:8]
    tasks_db[task_id] = {
        "task_id": task_id,
        "status": "pending",
        "created_at": datetime.now().isoformat(),
        "text": None,
        "summary": None
    }
    return task_id

def get_task(task_id: str) -> dict:
    """获取任务信息"""
    return tasks_db.get(task_id, {})

def update_task_status(task_id: str, status: str, error: str = None):
    """更新任务状态"""
    if task_id in tasks_db:
        tasks_db[task_id]["status"] = status
        if error:
            tasks_db[task_id]["error"] = error

def save_task_result(task_id: str, result: dict):
    """保存任务结果"""
    tasks_db[task_id].update(result)
    # 同时保存到文件
    result_file = EXPORT_DIR / f"{task_id}.json"
    with open(result_file, "w", encoding="utf-8") as f:
        json.dump(tasks_db[task_id], f, ensure_ascii=False, indent=2)

# ============ FastAPI 应用 ============

@asynccontextmanager
async def lifespan(app: FastAPI):
    """应用生命周期管理"""
    # 启动时预加载模型（可选）
    # asyncio.create_task(asyncio.to_thread(get_asr_model))
    yield
    # 清理
    tasks_db.clear()

app = FastAPI(
    title="语音转文字工具 API",
    description="支持 ASR 转写、LLM 重点提取、多格式导出",
    version="1.0.0",
    lifespan=lifespan
)

# CORS 配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ============ API 路由 ============

@app.post("/api/upload", response_model=dict)
async def upload_audio(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...)
):
    """上传音频文件并开始转写"""
    # 检查文件类型
    allowed_extensions = {'.mp3', '.wav', '.m4a', '.ogg', '.flac', '.aac', '.wma'}
    file_ext = Path(file.filename).suffix.lower()
    
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=400, 
            detail=f"不支持的文件格式: {file_ext}，请上传: {', '.join(allowed_extensions)}"
        )
    
    # 创建任务
    task_id = create_task()
    
    # 保存文件
    safe_filename = f"{task_id}{file_ext}"
    file_path = UPLOAD_DIR / safe_filename
    
    with open(file_path, "wb") as f:
        content = await file.read()
        f.write(content)
    
    # 后台处理
    background_tasks.add_task(process_audio_async, str(file_path), task_id)
    
    return {
        "task_id": task_id,
        "status": "pending",
        "message": "文件上传成功，正在转写中..."
    }

@app.get("/api/status/{task_id}", response_model=dict)
async def get_status(task_id: str):
    """获取任务状态"""
    task = get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    return task

@app.post("/api/summary", response_model=dict)
async def create_summary(request: SummaryRequest):
    """生成内容摘要"""
    task = get_task(request.task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    if task.get("status") != "completed":
        raise HTTPException(status_code=400, detail="转写尚未完成")
    
    text = task.get("text", "")
    if not text:
        raise HTTPException(status_code=400, detail="转写内容为空")
    
    # 生成摘要
    summary = await generate_summary(text, request.summary_type)
    
    # 保存摘要
    tasks_db[request.task_id]["summary"] = summary
    tasks_db[request.task_id]["summary_type"] = request.summary_type
    
    return {
        "task_id": request.task_id,
        "summary": summary,
        "summary_type": request.summary_type
    }

@app.post("/api/export")
async def export_result(request: ExportRequest):
    """导出转写结果"""
    task = get_task(request.task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    if task.get("status") != "completed":
        raise HTTPException(status_code=400, detail="转写尚未完成")
    
    format_handlers = {
        "txt": lambda: export_to_txt(task, request.include_summary),
        "srt": lambda: export_to_srt(task),
        "json": lambda: json.dumps(task, ensure_ascii=False, indent=2),
    }
    
    content_type_map = {
        "txt": "text/plain; charset=utf-8",
        "srt": "text/plain; charset=utf-8",
        "json": "application/json; charset=utf-8",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }
    
    if request.format == "docx":
        file_path = export_to_docx(task, request.include_summary)
        return FileResponse(
            file_path,
            media_type=content_type_map["docx"],
            filename=f"transcription_{request.task_id}.docx"
        )
    
    if request.format not in format_handlers:
        raise HTTPException(status_code=400, detail=f"不支持的导出格式: {request.format}")
    
    content = format_handlers[request.format]()
    
    # 保存到文件
    output_path = EXPORT_DIR / f"{request.task_id}.{request.format}"
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)
    
    return FileResponse(
        output_path,
        media_type=content_type_map[request.format],
        filename=f"transcription_{request.task_id}.{request.format}"
    )

@app.get("/api/tasks")
async def list_tasks():
    """获取所有任务列表"""
    return {
        "tasks": list(tasks_db.values())
    }

@app.delete("/api/tasks/{task_id}")
async def delete_task(task_id: str):
    """删除任务"""
    if task_id in tasks_db:
        del tasks_db[task_id]
    return {"message": "任务已删除"}

class SaveRequest(BaseModel):
    text: str

@app.post("/api/save/{task_id}")
async def save_transcription(task_id: str, request: SaveRequest):
    """保存编辑后的转写内容"""
    if task_id not in tasks_db:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    tasks_db[task_id]["text"] = request.text
    
    # 更新保存的文件
    result_file = EXPORT_DIR / f"{task_id}.json"
    with open(result_file, "w", encoding="utf-8") as f:
        json.dump(tasks_db[task_id], f, ensure_ascii=False, indent=2)
    
    return {"message": "保存成功", "task_id": task_id}

@app.get("/health")
async def health_check():
    """健康检查"""
    return {"status": "healthy", "service": "voice-to-text-api"}

# ============ 主入口 ============

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
